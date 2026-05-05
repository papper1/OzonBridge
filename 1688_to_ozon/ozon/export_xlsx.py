"""Export mapped Ozon products to Excel using template-driven columns."""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import logging
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from mapper.template_option_mapper import is_template_option_field, map_field_value
from mapper.template_loader import load_template
from translator.field_translator import convert_international_size_to_russian
from config import RESOURCE_ROOT


LOGGER = logging.getLogger("ozon.export_xlsx")
DEFAULT_RUSSIAN_SIZE = "46"

LEGACY_HEADER_ALIASES: dict[str, str] = {
    "Название товара": "Product name",
    "Бренд": "Brand",
    "Модель": "Model",
    "Цвет": "Product color",
    "Вес": "Weight in package, g",
    "Вес с упаковкой": "Weight in package, g",
    "Размер обуви": "Shoe size",
}

OZON_NUMBER_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "Package length, mm": ("Package length, mm", "Package length, mm*"),
    "Package width, mm": ("Package width, mm", "Package width, mm*"),
    "Package height, mm": ("Package height, mm", "Package height, mm*"),
    "Weight in package, g": ("Weight in package, g", "Weight in package, g*"),
}


def _extract_storage_capacity_gb(value: Any) -> str:
    text = str(value or "").upper()
    match = re.search(r"(\d+(?:\.\d+)?)\s*(TB|GB)", text)
    if not match:
        return ""

    amount = float(match.group(1))
    if match.group(2) == "TB":
        amount *= 1024
    return str(int(amount)) if amount.is_integer() else str(amount)


def _extract_first_numeric(value: Any) -> str:
    match = re.search(r"(\d+(?:\.\d+)?)", str(value or "").replace(",", "."))
    return match.group(1) if match else ""


def _format_ram_value(value: Any) -> str:
    text = str(value or "").strip().upper()
    match = re.search(r"(\d+(?:\.\d+)?)\s*GB\b", text)
    if not match:
        return str(value or "").strip()
    amount = match.group(1)
    amount = str(int(float(amount))) if float(amount).is_integer() else amount
    return f"{amount} GB"


def normalize_ozon_number(value: Any) -> int | None:
    """Convert text-like numeric input into an integer Excel number for Ozon."""
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value != value:
            return None
        return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))

    text = str(value).strip()
    if not text:
        return None

    # Excel often stores text numbers with a leading apostrophe.
    text = text.lstrip("'").strip()
    if not text:
        return None

    match = re.search(r"-?\d+(?:[.,]\d+)?", text.replace(" ", ""))
    if not match:
        return None

    candidate = match.group(0).replace(",", ".")
    try:
        number = Decimal(candidate)
    except InvalidOperation:
        return None

    if number.is_nan():
        return None
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _normalize_row_numeric_fields(row_data: dict[str, Any]) -> dict[str, Any]:
    normalized = {
        LEGACY_HEADER_ALIASES.get(str(header), str(header)): value
        for header, value in dict(row_data).items()
    }
    for aliases in OZON_NUMBER_HEADER_ALIASES.values():
        source_value = None
        for header in aliases:
            if header in normalized:
                source_value = normalized.get(header)
                break
        if source_value is None:
            continue

        parsed = normalize_ozon_number(source_value)
        for header in aliases:
            if header in normalized:
                normalized[header] = parsed if parsed is not None else normalized.get(header)
    return normalized


def _validate_ozon_required_numbers(row_data: dict[str, Any], article_code: str) -> list[str]:
    errors: list[str] = []
    for canonical_name, aliases in OZON_NUMBER_HEADER_ALIASES.items():
        value = None
        for header in aliases:
            if header in row_data:
                value = row_data.get(header)
                break

        parsed = normalize_ozon_number(value)
        if parsed is None:
            errors.append(canonical_name)
            continue

        for header in aliases:
            if header in row_data:
                row_data[header] = parsed
    if errors:
        LOGGER.warning(
            "Skipping Ozon numeric validation for article=%s, missing/invalid fields: %s",
            article_code or "<unknown>",
            ", ".join(errors),
        )
    return errors


def _expand_product_variants(products: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Expand mapped products so each hardware variant can be exported as its own row."""
    expanded: list[dict[str, Any]] = []

    for product in products:
        if not isinstance(product, dict):
            continue

        row_data = product.get("row_data")
        variants = product.get("variants")
        if not isinstance(row_data, dict) or not isinstance(variants, list) or len(variants) <= 1:
            if isinstance(row_data, dict):
                product = dict(product)
                product["row_data"] = _normalize_row_numeric_fields(row_data)
            expanded.append(product)
            continue

        for variant in variants:
            if not isinstance(variant, dict):
                continue

            variant_row = _normalize_row_numeric_fields(dict(row_data))
            template_name = str(product.get("template_name") or "").strip().lower()
            ram_value = str(variant.get("ram") or "").strip()
            storage_value = str(variant.get("storage") or "").strip()
            color_value = str(variant.get("color") or "").strip().lower()
            color_name_value = str(variant.get("color_name") or color_value).strip()
            sku_value = str(variant.get("sku") or "").strip()
            price_value = "249" if template_name == "hoodie" else _extract_first_numeric(variant.get("price"))
            variant_image_url = str(variant.get("image_url") or "").strip()
            product_images = [str(url).strip() for url in product.get("images", []) or [] if str(url).strip()]

            if "Random access memory" in variant_row and ram_value:
                variant_row["Random access memory"] = _format_ram_value(ram_value)
            if "Total SSD capacity, GB" in variant_row and storage_value:
                variant_row["Total SSD capacity, GB"] = _extract_storage_capacity_gb(storage_value)
            if "Product color" in variant_row and color_value:
                variant_row["Product color"] = (
                    map_field_value("Product color", color_value, product=product)
                    if is_template_option_field("Product color")
                    else color_value
                )
            if "Product color*" in variant_row and color_value:
                variant_row["Product color*"] = (
                    map_field_value("Product color*", color_value, product=product)
                    if is_template_option_field("Product color*")
                    else color_value
                )
            if "Color name" in variant_row and color_name_value:
                variant_row["Color name"] = color_name_value
            size_value = str(variant.get("size") or "").strip()
            russian_size_value = str(convert_international_size_to_russian(size_value).get("russian_size") or "").strip() or DEFAULT_RUSSIAN_SIZE
            if "Russian size*" in variant_row:
                variant_row["Russian size*"] = russian_size_value
            if "Manufacturer size" in variant_row and size_value:
                variant_row["Manufacturer size"] = size_value
            if "Article code*" in variant_row and sku_value:
                variant_row["Article code*"] = sku_value
            if "Price, CNY*" in variant_row and price_value:
                variant_row["Price, CNY*"] = price_value
            if "Link to the main image*" in variant_row and variant_image_url:
                variant_row["Link to the main image*"] = variant_image_url
            if "Links to additional photos" in variant_row:
                additional_images = [url for url in product_images if url != variant_image_url] if variant_image_url else product_images[1:]
                variant_row["Links to additional photos"] = "\n".join(additional_images)

            expanded_product = dict(product)
            expanded_product["row_data"] = variant_row
            expanded_product["variant"] = dict(variant)
            expanded_product["variants"] = [dict(variant)]
            expanded.append(expanded_product)

    return expanded


def _group_products_by_template(products: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for product in products:
        if not isinstance(product, dict):
            continue
        template_name = str(product.get("template_name") or "generic")
        grouped[template_name].append(product)
    return grouped


def _apply_merge_on_one_pdp(products: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Assign one shared numeric group per source URL for Merge on One PDP."""
    group_ids: dict[str, str] = {}
    next_group = 1
    normalized_products: list[dict[str, Any]] = []

    for product in products:
        if not isinstance(product, dict):
            normalized_products.append(product)
            continue

        row_data = product.get("row_data")
        if not isinstance(row_data, dict):
            normalized_products.append(product)
            continue

        source_url = str(product.get("source_url") or "").strip()
        group_key = source_url or str(row_data.get("Article code*") or product.get("name") or "").strip()
        if group_key not in group_ids:
            group_ids[group_key] = str(next_group)
            next_group += 1

        updated_product = dict(product)
        updated_row = dict(row_data)
        if "Merge on One PDP*" in updated_row:
            updated_row["Merge on One PDP*"] = group_ids[group_key]
        updated_product["row_data"] = updated_row
        normalized_products.append(updated_product)

    return normalized_products


def _resolve_headers(products: list[dict[str, Any]]) -> list[str]:
    headers: list[str] = []
    for product in products:
        row_data = product.get("row_data")
        if not isinstance(row_data, dict):
            continue
        for header in row_data.keys():
            if header not in headers:
                headers.append(header)
    return headers


def _append_simple_sheet(workbook: Workbook, sheet_name: str, products: list[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title=sheet_name[:31] or "Sheet")
    headers = _resolve_headers(products)
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for product in products:
        row_data = product.get("row_data")
        if not isinstance(row_data, dict):
            continue
        normalized_row = _normalize_row_numeric_fields(row_data)
        worksheet.append([normalized_row.get(header, "") for header in headers])

    worksheet.freeze_panes = "A2"


def _build_header_index(worksheet, header_row: int) -> dict[str, int]:
    header_index: dict[str, int] = {}
    for column in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=header_row, column=column).value
        if value:
            header_index[str(value).strip()] = column
    return header_index


def _resolve_template_column(header_index: dict[str, int], header_name: str) -> int | None:
    aliases = OZON_NUMBER_HEADER_ALIASES.get(header_name, (header_name,))
    for alias in aliases:
        column = header_index.get(alias)
        if column is not None:
            return column
    return None


def _build_links_output_path(output_file: Path) -> Path:
    """Build the companion workbook path that stores 1688 source links."""
    return output_file.with_name(f"{output_file.stem}_1688_links{output_file.suffix}")


def _export_source_links(products: list[dict[str, Any]], output_file: Path) -> str:
    """Export a lightweight workbook containing the 1688 source URLs."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "1688 Sources"

    headers = [
        "template_name",
        "product_name",
        "article_code",
        "supplier",
        "source_url",
    ]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for product in products:
        if not isinstance(product, dict):
            continue
        row_data = product.get("row_data") if isinstance(product.get("row_data"), dict) else {}
        product_name = (
            row_data.get("Product name")
            or row_data.get("Название товара")
            or product.get("name")
            or ""
        )
        article_code = row_data.get("Article code*") or ""
        worksheet.append(
            [
                str(product.get("template_name") or ""),
                str(product_name or ""),
                str(article_code or ""),
                str(product.get("supplier") or ""),
                str(product.get("source_url") or ""),
            ]
        )

    worksheet.freeze_panes = "A2"
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    return str(output_file)


def _export_with_real_template(
    products: list[dict[str, Any]],
    output_file: Path,
    template_config: dict[str, Any],
) -> str:
    workbook_template = Path(str(template_config.get("workbook_template") or "")).expanduser()
    if workbook_template and not workbook_template.is_absolute():
        workbook_template = RESOURCE_ROOT / workbook_template
    workbook = load_workbook(workbook_template)
    sheet_name = str(template_config.get("sheet_name") or workbook.sheetnames[0])
    header_row = int(template_config.get("header_row") or 1)
    data_start_row = int(template_config.get("data_start_row") or header_row + 1)
    worksheet = workbook[sheet_name]
    header_index = _build_header_index(worksheet, header_row)

    current_row = data_start_row
    for product in products:
        row_data = product.get("row_data")
        if not isinstance(row_data, dict):
            continue

        normalized_row = _normalize_row_numeric_fields(row_data)
        article_code = str(normalized_row.get("Article code*") or product.get("name") or "").strip()
        _validate_ozon_required_numbers(normalized_row, article_code)

        for header, value in normalized_row.items():
            column = _resolve_template_column(header_index, header)
            if column is None:
                continue
            worksheet.cell(row=current_row, column=column).value = value
        current_row += 1

    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    return str(output_file)


def export_to_xlsx(products: list[dict], output_path: str) -> str:
    """Export mapped Ozon products into an official workbook when configured."""
    output_file = Path(output_path)
    expanded_products = _apply_merge_on_one_pdp(_expand_product_variants(products))
    grouped_products = _group_products_by_template(expanded_products)
    all_products = [product for template_products in grouped_products.values() for product in template_products]

    if len(grouped_products) == 1:
        template_name, template_products = next(iter(grouped_products.items()))
        template_config = load_template(template_name)
        workbook_template = Path(str(template_config.get("workbook_template") or "")).expanduser()
        if workbook_template and not workbook_template.is_absolute():
            workbook_template = RESOURCE_ROOT / workbook_template
        if workbook_template.is_file():
            exported_path = _export_with_real_template(template_products, output_file, template_config)
            _export_source_links(all_products, _build_links_output_path(output_file))
            return exported_path

    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for template_name, template_products in grouped_products.items():
        _append_simple_sheet(workbook, template_name, template_products)

    if not grouped_products:
        workbook.create_sheet(title="generic")

    workbook.save(output_file)
    _export_source_links(all_products, _build_links_output_path(output_file))
    return str(output_file)


def export_products(products: list[dict], output_path: str) -> str:
    """Compatibility wrapper for existing pipeline code."""
    return export_to_xlsx(products, output_path)
