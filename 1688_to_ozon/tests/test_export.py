import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from ozon.export_xlsx import export_to_xlsx, normalize_ozon_number


class TestExport(unittest.TestCase):
    def test_normalize_ozon_number_rounds_and_strips_text_markers(self) -> None:
        self.assertEqual(normalize_ozon_number("'357.6"), 358)
        self.assertEqual(normalize_ozon_number(" 229.7 "), 230)
        self.assertEqual(normalize_ozon_number("19.2"), 19)
        self.assertEqual(normalize_ozon_number("1500"), 1500)
        self.assertIsNone(normalize_ozon_number(""))

    def test_export_groups_products_by_template(self) -> None:
        products = [
            {
                "template_name": "laptop",
                "row_data": {
                    "Product name": "Lenovo ThinkBook laptop",
                    "Brand": "Lenovo",
                },
            },
            {
                "template_name": "clothing",
                "row_data": {
                    "Product name": "Cotton T-Shirt",
                    "Brand": "OEM",
                    "Product color": "Black",
                },
            },
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "export.xlsx"
            links_output_path = Path(temp_dir) / "export_1688_links.xlsx"
            export_to_xlsx(products, str(output_path))

            self.assertTrue(output_path.exists())
            self.assertTrue(links_output_path.exists())
            workbook = load_workbook(output_path)
            self.assertIn("laptop", workbook.sheetnames)
            self.assertIn("clothing", workbook.sheetnames)
            self.assertEqual(workbook["laptop"]["A2"].value, "Lenovo ThinkBook laptop")

            links_workbook = load_workbook(links_output_path)
            self.assertEqual(links_workbook["1688 Sources"]["A2"].value, "laptop")

    def test_export_normalizes_legacy_headers_to_english(self) -> None:
        products = [
            {
                "template_name": "generic",
                "row_data": {
                    "Название товара": "Legacy Product",
                    "Бренд": "OEM",
                    "Цвет": "Black",
                },
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "legacy.xlsx"
            export_to_xlsx(products, str(output_path))

            workbook = load_workbook(output_path)
            worksheet = workbook["generic"]

            self.assertEqual(worksheet["A1"].value, "Product name")
            self.assertEqual(worksheet["B1"].value, "Brand")
            self.assertEqual(worksheet["C1"].value, "Product color")
            self.assertEqual(worksheet["A2"].value, "Legacy Product")

    def test_export_expands_variant_rows(self) -> None:
        products = [
            {
                "template_name": "generic",
                "row_data": {
                    "Article code*": "base-sku",
                    "Product name": "Laptop M1",
                    "Price, CNY*": "100",
                    "Random access memory": "",
                    "Product color": "",
                    "Color name": "",
                    "Russian size*": "",
                    "Manufacturer size": "",
                    "Total SSD capacity, GB": "",
                    "Link to the main image*": "https://img.example.com/base.jpg",
                    "Links to additional photos": "https://img.example.com/detail.jpg",
                },
                "images": [
                    "https://img.example.com/base.jpg",
                    "https://img.example.com/detail.jpg",
                ],
                "variants": [
                    {"sku": "m1-6gb-128gb", "ram": "6GB", "storage": "128GB", "color": "Grey", "color_name": "Space Grey", "size": "XL", "price": "100", "image_url": "https://img.example.com/grey.jpg"},
                    {"sku": "m1-6gb-256gb", "ram": "6GB", "storage": "256GB", "color": "Silver", "color_name": "Moonlight Silver", "price": "120", "image_url": "https://img.example.com/silver.jpg"},
                ],
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "variants.xlsx"
            export_to_xlsx(products, str(output_path))

            workbook = load_workbook(output_path)
            worksheet = workbook["generic"]

            self.assertEqual(worksheet["A2"].value, "m1-6gb-128gb")
            self.assertEqual(worksheet["D2"].value, "6 GB")
            self.assertEqual(worksheet["E2"].value, "grey")
            self.assertEqual(worksheet["F2"].value, "Space Grey")
            self.assertEqual(worksheet["G2"].value, "50")
            self.assertEqual(worksheet["H2"].value, "XL")
            self.assertEqual(worksheet["I2"].value, "128")
            self.assertEqual(worksheet["J2"].value, "https://img.example.com/grey.jpg")
            self.assertEqual(worksheet["K2"].value, "https://img.example.com/base.jpg\nhttps://img.example.com/detail.jpg")
            self.assertEqual(worksheet["A3"].value, "m1-6gb-256gb")
            self.assertEqual(worksheet["D3"].value, "6 GB")
            self.assertEqual(worksheet["E3"].value, "silver")
            self.assertEqual(worksheet["F3"].value, "Moonlight Silver")
            self.assertEqual(worksheet["I3"].value, "256")
            self.assertEqual(worksheet["J3"].value, "https://img.example.com/silver.jpg")

    def test_export_writes_package_fields_as_real_numbers_in_template(self) -> None:
        products = [
            {
                "template_name": "laptop",
                "row_data": {
                    "Article code*": "sku-1",
                    "Package width, mm*": "'357.6",
                    "Package height, mm*": "229.7",
                    "Package length, mm*": "19.2",
                    "Weight in package, g*": "1500",
                },
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "template.xlsx"
            output_path = Path(temp_dir) / "export.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Template"
            worksheet["A2"] = "Article code*"
            worksheet["B2"] = "Package length, mm*"
            worksheet["C2"] = "Package width, mm*"
            worksheet["D2"] = "Package height, mm*"
            worksheet["E2"] = "Weight in package, g*"
            workbook.save(template_path)

            template_config = {
                "template_name": "laptop",
                "workbook_template": str(template_path),
                "sheet_name": "Template",
                "header_row": 2,
                "data_start_row": 5,
            }

            with patch("ozon.export_xlsx.load_template", return_value=template_config):
                export_to_xlsx(products, str(output_path))

            exported = load_workbook(output_path)
            worksheet = exported["Template"]

            self.assertEqual(worksheet["A5"].value, "sku-1")
            self.assertEqual(worksheet["B5"].value, 19)
            self.assertEqual(worksheet["C5"].value, 358)
            self.assertEqual(worksheet["D5"].value, 230)
            self.assertEqual(worksheet["E5"].value, 1500)
            self.assertEqual(worksheet["B5"].data_type, "n")
            self.assertEqual(worksheet["C5"].data_type, "n")
            self.assertEqual(worksheet["D5"].data_type, "n")
            self.assertEqual(worksheet["E5"].data_type, "n")

    def test_export_forces_hoodie_variant_price_to_249(self) -> None:
        products = [
            {
                "template_name": "hoodie",
                "source_url": "https://detail.1688.com/offer/987041511404.html",
                "row_data": {
                    "Article code*": "hoodie-base",
                    "Product name": "Oversized Hoodie",
                    "Price, CNY*": "249",
                    "Merge on One PDP*": "",
                    "Product color*": "",
                    "Color name": "",
                    "Russian size*": "",
                    "Height": "",
                },
                "variants": [
                    {"sku": "hoodie-black-m", "color": "Black", "color_name": "Black", "size": "M", "price": "89"},
                    {"sku": "hoodie-black-l", "color": "Black", "color_name": "Black", "size": "L", "price": "99"},
                ],
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "hoodie.xlsx"
            with patch(
                "ozon.export_xlsx.load_template",
                return_value={"template_name": "hoodie", "workbook_template": str(Path(temp_dir) / "missing.xlsx")},
            ):
                export_to_xlsx(products, str(output_path))

            workbook = load_workbook(output_path)
            worksheet = workbook["hoodie"]

            self.assertEqual(worksheet["A2"].value, "hoodie-black-m")
            self.assertEqual(worksheet["C2"].value, "249")
            self.assertEqual(worksheet["G2"].value, "46")
            self.assertIsNone(worksheet["H2"].value)
            self.assertEqual(worksheet["A3"].value, "hoodie-black-l")
            self.assertEqual(worksheet["C3"].value, "249")

    def test_export_maps_hoodie_product_color_to_allowed_template_value_but_keeps_color_name_freeform(self) -> None:
        products = [
            {
                "template_name": "hoodie",
                "source_url": "https://www.etsy.com/listing/4456307613/demo",
                "row_data": {
                    "Article code*": "hoodie-base",
                    "Product name": "Oversized Hoodie",
                    "Price, CNY*": "249",
                    "Merge on One PDP*": "",
                    "Product color*": "",
                    "Color name": "",
                    "Russian size*": "",
                    "Height": "",
                },
                "variants": [
                    {
                        "sku": "hoodie-baby-pink-m",
                        "color": "baby pink hoodie ($67.83)",
                        "color_name": "baby pink hoodie ($67.83)",
                        "size": "M",
                        "price": "67.83",
                    },
                ],
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "hoodie_colors.xlsx"
            with patch(
                "ozon.export_xlsx.load_template",
                return_value={"template_name": "hoodie", "workbook_template": str(Path(temp_dir) / "missing.xlsx")},
            ):
                export_to_xlsx(products, str(output_path))

            workbook = load_workbook(output_path)
            worksheet = workbook["hoodie"]

            self.assertEqual(worksheet["E2"].value, "Baby Pink")
            self.assertEqual(worksheet["F2"].value, "baby pink hoodie ($67.83)")

    def test_export_assigns_same_merge_on_one_pdp_for_variants_from_same_link(self) -> None:
        products = [
            {
                "template_name": "hoodie",
                "source_url": "https://detail.1688.com/offer/111.html",
                "row_data": {
                    "Article code*": "hoodie-base-1",
                    "Merge on One PDP*": "",
                    "Product name": "Hoodie A",
                },
                "variants": [
                    {"sku": "hoodie-a-black-m", "color": "Black", "color_name": "Black", "size": "M", "price": "89"},
                    {"sku": "hoodie-a-white-l", "color": "White", "color_name": "White", "size": "L", "price": "99"},
                ],
            },
            {
                "template_name": "hoodie",
                "source_url": "https://detail.1688.com/offer/222.html",
                "row_data": {
                    "Article code*": "hoodie-base-2",
                    "Merge on One PDP*": "",
                    "Product name": "Hoodie B",
                },
            },
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "merge_pdp.xlsx"
            with patch(
                "ozon.export_xlsx.load_template",
                return_value={"template_name": "hoodie", "workbook_template": str(Path(temp_dir) / "missing.xlsx")},
            ):
                export_to_xlsx(products, str(output_path))

            workbook = load_workbook(output_path)
            worksheet = workbook["hoodie"]

            self.assertEqual(worksheet["B2"].value, "1")
            self.assertEqual(worksheet["B3"].value, "1")
            self.assertEqual(worksheet["B4"].value, "2")


if __name__ == "__main__":
    unittest.main()
